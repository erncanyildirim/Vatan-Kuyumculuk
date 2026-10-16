"""Writes store-assets/metinler/yetki-belgesi.docx, the printable form of yetki-belgesi.md.

usage: python3 tools/store/make_authorization_letter.py   (run from the project root; needs python-docx)
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

BLANK = "………………………………………………"

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.2)
    s.left_margin = s.right_margin = Cm(2.4)
style = doc.styles["Normal"]
style.font.name = "Arial"
style.font.size = Pt(11)


def para(text="", bold=False, size=None, align=None, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def field(label, value=BLANK):
    p = doc.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    p.add_run(value)
    p.paragraph_format.space_after = Pt(4)


para("YETKİ BELGESİ / AUTHORIZATION LETTER", bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
field("Tarih / Date", "…… / …… / 2026")
field("İşletme / Business", "Vatan Kuyumculuk")
field("İşletme sahibi / Owner")
field("Adres / Address", BLANK + " Osmaniye, Türkiye")
field("Vergi dairesi ve no / Tax office and ID")
field("Telefon / Phone", "+90 505 434 85 92")
para(space_after=8)
para(
    "Vatan Kuyumculuk'un sahibi olarak, Eren Can Yıldırım adlı kişiye; işletmemizin adını, logosunu ve "
    "iletişim bilgilerini içeren \"Vatan Kuyumculuk Osmaniye\" mobil uygulamasını kendi Apple Developer ve "
    "Google Play geliştirici hesapları üzerinden yayınlama, güncelleme ve yönetme yetkisi verdiğimi beyan ederim.",
    space_after=12,
)
para(
    "As the owner of Vatan Kuyumculuk, I hereby authorize Eren Can Yildirim to publish, update and manage the "
    "\"Vatan Kuyumculuk Osmaniye\" mobile application, which uses our business name, logo and contact "
    "information, through their personal Apple Developer and Google Play developer accounts.",
    space_after=12,
)
field("Uygulama kimliği / App identifier", "com.suer.vatankuyumculuk")
para(space_after=24)
field("Ad Soyad / Name")
para(space_after=18)
field("İmza ve kaşe / Signature and stamp", "")
doc.save("store-assets/metinler/yetki-belgesi.docx")
print("wrote store-assets/metinler/yetki-belgesi.docx")
